"""Policy detail import: explicit merges, summaries, dates and text identities."""
import io
import zipfile

import pytest
import pandas as pd
from openpyxl import Workbook

from etl.parser import parse_jingdai_excel


HEADERS = ['统计时间_年月日', '销售平台', '投保单号', '产品代码', '产品名称',
           '产品销售名称', '长短险', '缴别', '是否社会保障类产品',
           '是否商保年金产品', '是否个人养老金', '缴费年限', '期交保费', '年化规保']


def workbook(rows, *, extra_filter=False, merges=(), date1904=False):
    book = Workbook()
    sheet = book.active
    sheet.append(['统计时间_年', '等于', '2026'])
    if extra_filter:
        sheet.append(['统计时间_年月', '小于等于', '2026-06'])
    sheet.append([])
    sheet.append(HEADERS)
    for row in rows:
        sheet.append(row)
    for merge in merges:
        sheet.merge_cells(merge)
    if date1904:
        from openpyxl.utils.datetime import CALENDAR_MAC_1904
        book.epoch = CALENDAR_MAC_1904
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


def record(**changes):
    values = {'统计时间_年月日': 46023, '销售平台': '合成平台', '投保单号': '000001',
              '产品代码': '0009', '产品名称': '合成产品', '产品销售名称': '合成名称',
              '长短险': '长险', '缴别': '月交', '是否社会保障类产品': 'Y',
              '是否商保年金产品': 'N', '是否个人养老金': 'N',
              '缴费年限': 5, '期交保费': 100, '年化规保': 120}
    values.update(changes)
    return [values[key] for key in HEADERS]


@pytest.mark.parametrize('extra_filter', [True, False])
def test_new_header_is_detected_and_canonical_aliases_preserve_identity(extra_filter):
    parsed = parse_jingdai_excel(workbook([record()], extra_filter=extra_filter))
    assert parsed['时间'].tolist() == ['2026-01-01']
    assert parsed['投保单号'].tolist() == ['000001']
    assert parsed['产品代码'].tolist() == ['0009']
    assert parsed['经代机构'].tolist() == parsed['销售平台'].tolist() == ['合成平台']
    assert parsed['承保年化规保'].tolist() == [120]
    assert parsed['经代源格式'].tolist() == ['policy_detail']
    assert parsed['承保件数'].tolist() == [1]
    assert parsed['件数来源'].tolist() == ['policy_identity_daily']
    assert parsed['缴费年限'].tolist() == [5]
    assert parsed['当前缴别大类'].tolist() == ['期交']


def test_only_explicit_merge_cells_are_restored_and_subtotals_are_removed():
    first = record()
    second = record(统计时间_年月日=None, 销售平台=None, 投保单号=None, 期交保费=-20, 年化规保=-24)
    subtotal = [None, '小计', None, None, None, None, None, None, None, None, None, None, 80, 96]
    blank_identity = record(投保单号=None, 期交保费=10, 年化规保=12)
    parsed = parse_jingdai_excel(workbook([first, second, subtotal, blank_identity],
        merges=['A4:A6', 'B4:B5', 'C4:C5']))
    assert len(parsed) == 3
    assert parsed['投保单号'].iloc[:2].tolist() == ['000001', '000001']
    assert pd.isna(parsed['投保单号'].iloc[2])
    assert parsed['时间'].tolist() == ['2026-01-01'] * 3
    assert parsed['期交保费'].sum() == 90
    assert parsed.attrs['ignored_summary_rows'] == 1
    assert parsed.attrs['restored_merge_cells'] == 4
    assert parsed['承保件数'].iloc[:2].tolist() == [1, 0]
    assert pd.isna(parsed['承保件数'].iloc[2])


def test_duplicates_zero_and_negative_transactions_are_kept():
    parsed = parse_jingdai_excel(workbook([record(), record(),
        record(期交保费=0, 年化规保=0), record(期交保费=-100, 年化规保=-120)]))
    assert len(parsed) == 4
    assert parsed['期交保费'].tolist() == [100, 100, 0, -100]
    assert parsed['投保单号'].nunique() == 1
    assert parsed['承保件数'].tolist() == [1, 0, 0, 0]


def test_invalid_missing_date_and_amount_are_not_silently_filled():
    parsed = parse_jingdai_excel(workbook([record(), record(统计时间_年月日=None,
        销售平台=None, 投保单号=None, 期交保费='bad')]))
    assert pd.isna(parsed.iloc[1]['时间'])
    assert pd.isna(parsed.iloc[1]['经代机构'])
    assert pd.isna(parsed.iloc[1]['投保单号'])
    assert parsed.iloc[1]['期交保费'] == 'bad'


def test_numeric_and_yyyymmdd_dates_and_1904_epoch_are_supported():
    parsed = parse_jingdai_excel(workbook([record(统计时间_年月日='20260102')]))
    assert parsed['时间'].tolist() == ['2026-01-02']
    parsed = parse_jingdai_excel(workbook([record(统计时间_年月日=44562)], date1904=True))
    assert parsed['时间'].tolist() == ['2026-01-02']


def test_windows_archive_paths_remain_readable():
    original = workbook([record()])
    stream = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(original)) as source, zipfile.ZipFile(stream, 'w') as target:
        for name in source.namelist():
            target.writestr(name.replace('/', '\\'), source.read(name))
    assert parse_jingdai_excel(stream.getvalue())['投保单号'].tolist() == ['000001']

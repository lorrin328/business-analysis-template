import os
import sys
from io import BytesIO

from openpyxl import load_workbook

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
sys.path.insert(0, os.path.join(ROOT, "backend"))


def test_honor_export_builds_workbook(auth_db):
    from honor.repository import create_batch, replace_calculation_results
    from honor.exporter import build_honor_export_workbook

    batch_id = create_batch(year=2026, month=5, rule_version="2026-v1", created_by="pytest")
    replace_calculation_results(
        batch_id,
        {
            "org_summary": [{"batch_id": batch_id, "year": 2026, "month": 5, "org": "上海", "business_line": "OTO", "tracked_headcount": 1}],
            "person_summary": [{"batch_id": batch_id, "year": 2026, "latest_month": 5, "org": "上海", "business_line": "OTO", "staff_code": "1001", "membership_level": "初级会员"}],
            "person_month": [{"batch_id": batch_id, "year": 2026, "month": 5, "org": "上海", "business_line": "OTO", "staff_code": "1001", "membership_level": "初级会员"}],
            "quarter_rewards": [],
            "exceptions": [],
            "source_staff_month": [],
            "source_policy": [],
        },
        0,
    )
    content = build_honor_export_workbook(batch_id)
    assert content.startswith(b"PK")
    assert len(content) > 1000


def test_honor_export_does_not_truncate_after_5000_rows(auth_db):
    from honor.repository import create_batch, replace_calculation_results
    from honor.exporter import build_honor_export_workbook

    batch_id = create_batch(year=2026, month=6, rule_version="2026-v1", created_by="pytest")
    rows = [{"batch_id": batch_id, "year": 2026, "month": 6, "org": "上海",
             "business_line": "OTO", "staff_code": f"test-{index}", "membership_level": "初级会员"}
            for index in range(5001)]
    replace_calculation_results(batch_id, {
        "org_summary": [], "person_summary": [], "person_month": rows,
        "quarter_rewards": [], "exceptions": [], "source_staff_month": [], "source_policy": [],
    }, 0)
    workbook = load_workbook(BytesIO(build_honor_export_workbook(batch_id)), read_only=True)
    try:
        assert workbook["月度明细"].max_row == 5002
        assert list(workbook["导出校验"].values)[3] == ("月度明细", 5001, 5001, "一致")
    finally:
        workbook.close()
